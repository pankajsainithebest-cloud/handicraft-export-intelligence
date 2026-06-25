"""Ingestion engine tests with fixture workbooks."""

from pathlib import Path

from openpyxl import Workbook, load_workbook

from handicraft_export_intelligence.ingestion.engine import run_ingestion
from handicraft_export_intelligence.settings import load_settings


def test_run_ingestion_generates_outputs_and_continues_after_errors(
    tmp_path: Path,
) -> None:
    data_dir = tmp_path / "data"
    output_dir = tmp_path / "output"
    data_dir.mkdir()
    output_dir.mkdir()
    _create_fixture_workbook(data_dir / "Ajmer.xlsx")
    (data_dir / "Broken.xlsx").write_text("not an excel workbook", encoding="utf-8")

    settings = load_settings()
    test_settings = settings.model_copy(
        update={
            "paths": settings.paths.model_copy(
                update={
                    "raw_workbook_dir": data_dir,
                    "output_dir": output_dir,
                }
            )
        }
    )

    result = run_ingestion(test_settings)

    assert len(result.workbook_inventory) == 2
    assert result.successful_workbooks == 1
    assert result.failed_workbooks == 1
    assert result.product_sheet_count == 1
    assert result.no_details_sheet_count == 1
    assert result.output_dir == output_dir / "ingestion"
    assert (result.output_dir / "workbook_inventory.xlsx").exists()
    assert (result.output_dir / "sheet_inventory.xlsx").exists()
    assert (result.output_dir / "ingestion_summary.json").exists()
    assert (result.output_dir / "ingestion_report.md").exists()

    inventory = load_workbook(
        result.output_dir / "sheet_inventory.xlsx",
        read_only=True,
        data_only=True,
    )
    try:
        rows = list(inventory.active.iter_rows(values_only=True))
    finally:
        inventory.close()

    assert rows[0][0] == "workbook"
    assert any(row[2] == "Blue Pottery" for row in rows[1:])


def _create_fixture_workbook(path: Path) -> None:
    workbook = Workbook()
    product_sheet = workbook.active
    product_sheet.title = "Blue Pottery"
    product_sheet.append(["ArtisanId", "Name", "Product", "District"])
    product_sheet.append(["A-1", "Meera", "Blue Pottery", "Ajmer"])

    no_details = workbook.create_sheet("No Details")
    no_details.append(["Message"])
    no_details.append(["No records"])

    workbook.save(path)
