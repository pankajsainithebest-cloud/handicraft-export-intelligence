"""Read source Excel workbooks in read-only mode."""

from __future__ import annotations

import logging
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook

from handicraft_export_intelligence.ingestion.models import (
    ParsedProductSheet,
    SheetMetadata,
    WorkbookMetadata,
)
from handicraft_export_intelligence.ingestion.sheet_parser import (
    classify_sheet,
    metadata_for_non_product_sheet,
    parse_product_sheet,
)
from handicraft_export_intelligence.ingestion.workbook_discovery import (
    district_from_workbook,
)

logger = logging.getLogger(__name__)


def read_workbook(
    workbook_path: Path,
    no_details_sheet: str,
) -> tuple[WorkbookMetadata, tuple[SheetMetadata, ...], tuple[ParsedProductSheet, ...]]:
    """Read one workbook and return workbook, sheet, and raw parsed data metadata."""

    district = district_from_workbook(workbook_path)
    base_metadata = _workbook_metadata(workbook_path, district)
    logger.info("Processing workbook: %s", workbook_path)

    try:
        workbook = load_workbook(
            filename=workbook_path,
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        logger.exception("Failed to open workbook: %s", workbook_path)
        return (
            WorkbookMetadata(
                workbook=base_metadata.workbook,
                district=base_metadata.district,
                path=base_metadata.path,
                size_bytes=base_metadata.size_bytes,
                modified_at=base_metadata.modified_at,
                status="error",
                error=str(exc),
            ),
            (),
            (),
        )

    sheet_inventory: list[SheetMetadata] = []
    parsed_product_sheets: list[ParsedProductSheet] = []

    try:
        for worksheet in workbook.worksheets:
            logger.info(
                "Processing sheet: workbook=%s sheet=%s",
                workbook_path.name,
                worksheet.title,
            )
            sheet_type = classify_sheet(worksheet.title, no_details_sheet)
            try:
                if sheet_type == "product":
                    parsed_sheet = parse_product_sheet(
                        worksheet=worksheet,
                        workbook=workbook_path.name,
                        district=district,
                    )
                    sheet_inventory.append(parsed_sheet.metadata)
                    parsed_product_sheets.append(parsed_sheet)
                else:
                    sheet_inventory.append(
                        metadata_for_non_product_sheet(
                            worksheet=worksheet,
                            workbook=workbook_path.name,
                            district=district,
                            sheet_type=sheet_type,
                        )
                    )
            except Exception as exc:
                logger.exception(
                    "Failed to process sheet: workbook=%s sheet=%s",
                    workbook_path.name,
                    worksheet.title,
                )
                sheet_inventory.append(
                    SheetMetadata(
                        workbook=workbook_path.name,
                        district=district,
                        sheet=worksheet.title,
                        sheet_type=sheet_type,
                        row_count=worksheet.max_row,
                        column_count=worksheet.max_column,
                        header_count=0,
                        parsed_data_rows=0,
                        status="error",
                        error=str(exc),
                    )
                )
    finally:
        workbook.close()

    metadata = WorkbookMetadata(
        workbook=base_metadata.workbook,
        district=base_metadata.district,
        path=base_metadata.path,
        size_bytes=base_metadata.size_bytes,
        modified_at=base_metadata.modified_at,
        sheet_count=len(sheet_inventory),
        status="ok",
    )
    return metadata, tuple(sheet_inventory), tuple(parsed_product_sheets)


def _workbook_metadata(workbook_path: Path, district: str) -> WorkbookMetadata:
    stat = workbook_path.stat()
    modified_at = datetime.fromtimestamp(stat.st_mtime, tz=UTC).isoformat()
    return WorkbookMetadata(
        workbook=workbook_path.name,
        district=district,
        path=workbook_path,
        size_bytes=stat.st_size,
        modified_at=modified_at,
    )
