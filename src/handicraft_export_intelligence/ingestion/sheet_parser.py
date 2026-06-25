"""Raw worksheet parsing without cleaning or normalization."""

from __future__ import annotations

from collections.abc import Sequence
from typing import Any

from openpyxl.worksheet.worksheet import Worksheet

from handicraft_export_intelligence.ingestion.models import (
    ParsedProductSheet,
    SheetMetadata,
)


def classify_sheet(sheet_name: str, no_details_sheet: str) -> str:
    """Classify a worksheet as product or No Details."""

    if sheet_name.casefold() == no_details_sheet.casefold():
        return "no_details"
    return "product"


def parse_product_sheet(
    worksheet: Worksheet,
    workbook: str,
    district: str,
) -> ParsedProductSheet:
    """Parse a product worksheet while preserving raw cell values."""

    rows = worksheet.iter_rows(values_only=True)
    headers = _normalize_headers(next(rows, ()))
    parsed_rows: list[dict[str, Any]] = []

    for row_number, row_values in enumerate(rows, start=2):
        if _is_empty_row(row_values):
            continue
        record = {
            "_source_workbook": workbook,
            "_source_district": district,
            "_source_sheet": worksheet.title,
            "_source_row_number": row_number,
        }
        for index, header in enumerate(headers):
            record[header] = row_values[index] if index < len(row_values) else None
        parsed_rows.append(record)

    metadata = SheetMetadata(
        workbook=workbook,
        district=district,
        sheet=worksheet.title,
        sheet_type="product",
        row_count=worksheet.max_row,
        column_count=worksheet.max_column,
        header_count=len(headers),
        parsed_data_rows=len(parsed_rows),
    )
    return ParsedProductSheet(metadata=metadata, rows=tuple(parsed_rows))


def metadata_for_non_product_sheet(
    worksheet: Worksheet,
    workbook: str,
    district: str,
    sheet_type: str,
) -> SheetMetadata:
    """Capture metadata for a sheet that is not parsed as a product sheet."""

    headers = _normalize_headers(next(worksheet.iter_rows(values_only=True), ()))
    return SheetMetadata(
        workbook=workbook,
        district=district,
        sheet=worksheet.title,
        sheet_type=sheet_type,
        row_count=worksheet.max_row,
        column_count=worksheet.max_column,
        header_count=len(headers),
        parsed_data_rows=0,
    )


def _normalize_headers(header_values: Sequence[Any]) -> tuple[str, ...]:
    headers: list[str] = []
    for index, value in enumerate(header_values, start=1):
        if value is None or str(value).strip() == "":
            headers.append(f"Unnamed Column {index}")
        else:
            headers.append(str(value))
    return tuple(headers)


def _is_empty_row(row_values: Sequence[Any]) -> bool:
    return all(value is None for value in row_values)
